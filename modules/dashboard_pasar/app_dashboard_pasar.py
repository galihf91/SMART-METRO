import streamlit as st
import pandas as pd
import folium
from folium.plugins import MarkerCluster
from streamlit_folium import st_folium
from datetime import datetime
from io import StringIO, BytesIO
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
import re
import numpy as np
import json
import glob
import os
from supabase import create_client

st.set_page_config(
    page_title="Dashboard Pasar – SMART METRO",
    page_icon="🏪",
    layout="wide"
)

# =========================
# KONSTANTA DASHBOARD PASAR
# =========================
FILE_EXCEL = "data/DATA_DASHBOARD_PASAR.xlsx"
FILE_GEOJSON = "data/batas_kecamatan_tangerang.geojson"

def render_main_header(title, subtitle):
    html = (
        '<div style="'
        'background:linear-gradient(135deg,#4B0082 0%,#6D28D9 55%,#7C3AED 100%);'
        'padding:22px 26px;'
        'border-radius:16px;'
        'margin-bottom:22px;'
        'box-shadow:0 8px 20px rgba(75,0,130,0.18);'
        'color:white;'
        'display:flex;'
        'align-items:center;'
        'gap:16px;'
        '">'
        
        '<div style="'
        'width:52px;'
        'height:52px;'
        'border-radius:14px;'
        'background:rgba(255,255,255,0.15);'
        'display:flex;'
        'align-items:center;'
        'justify-content:center;'
        'font-size:28px;'
        'flex-shrink:0;'
        '">'
        '🏪'
        '</div>'
        
        '<div style="flex:1;">'
        
        '<div style="'
        'font-size:12px;'
        'font-weight:700;'
        'letter-spacing:0.8px;'
        'opacity:0.85;'
        'margin-bottom:3px;'
        '">'
        'SMART METRO · TERA ULANG PASAR'
        '</div>'
        
        '<div style="'
        'font-size:28px;'
        'font-weight:800;'
        'line-height:1.2;'
        'margin-bottom:5px;'
        '">'
        f'{title}'
        '</div>'
        
        '<div style="'
        'font-size:14px;'
        'opacity:0.9;'
        'line-height:1.4;'
        '">'
        f'{subtitle}'
        '</div>'
        
        '</div>'
        '</div>'
    )

    st.markdown(
        html,
        unsafe_allow_html=True
    )
# =========================
# FUNGSI UTILITAS UMUM
# =========================
def hitung_perubahan_persen(nilai_sekarang, nilai_sebelumnya):
    if nilai_sebelumnya is None or pd.isna(nilai_sebelumnya):
        return None

    if nilai_sebelumnya == 0:
        return None

    return (
        (nilai_sekarang - nilai_sebelumnya)
        / nilai_sebelumnya
        * 100
    )
def _norm(s): return re.sub(r"[^a-z0-9]+", "", str(s).strip().lower())

def parse_coord(val):
    try:
        if pd.isna(val) or val == "": return np.nan, np.nan
        s = str(val).strip()
        if ',' in s:
            lat, lon = map(float, s.split(',')[:2])
            if abs(lat) > 90: lat, lon = lon, lat
            return lat, lon
        nums = re.findall(r"-?\d+(?:\.\d+)?", s)
        if len(nums) >= 2:
            lat, lon = map(float, nums[:2])
            if abs(lat) > 90: lat, lon = lon, lat
            return lat, lon
    except: pass
    return np.nan, np.nan

def uniq(series, clean=False):
    s = series.dropna().astype(str).str.strip()
    if clean: s = s.str.title()
    s = s[~s.str.lower().isin(["", "nan", "none", "null", "na", "n/a", "-", "--"])]
    return sorted(s.unique())

def marker_color(year, selected_year):
    if year is None or year == 0: return "gray"
    if year == selected_year: return "green"
    if year == selected_year - 1: return "orange"
    return "red"


# =========================
# LOAD DATA PASAR
# =========================
@st.cache_data(ttl=60)
def load_master_pasar_supabase():
    sb = get_supabase()

    data = (
        sb.table("pasar")
        .select(
            "id,nama_pasar,alamat,kecamatan,latitude,longitude,status"
        )
        .execute()
        .data
        or []
    )

    df = pd.DataFrame(data)

    if df.empty:
        return df

    df["lat"] = pd.to_numeric(
        df["latitude"],
        errors="coerce"
    )

    df["lon"] = pd.to_numeric(
        df["longitude"],
        errors="coerce"
    )

    df["nama_pasar"] = (
        df["nama_pasar"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    df["kecamatan"] = (
        df["kecamatan"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    df["alamat"] = (
        df["alamat"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    return df
@st.cache_resource
def get_supabase():
    url = None
    key = None

    try:
        url = st.secrets.get("SUPABASE_URL")
        key = (
            st.secrets.get("SUPABASE_KEY")
            or st.secrets.get("SUPABASE_ANON_KEY")
        )
    except Exception:
        pass

    if not url or not key:
        try:
            cfg = st.secrets.get("supabase", {})
            url = url or cfg.get("url")
            key = (
                key
                or cfg.get("key")
                or cfg.get("anon_key")
            )
        except Exception:
            pass

    url = url or os.getenv("SUPABASE_URL")
    key = (
        key
        or os.getenv("SUPABASE_KEY")
        or os.getenv("SUPABASE_ANON_KEY")
    )

    if not url or not key:
        st.error(
            "Koneksi Supabase belum tersedia. "
            "Pastikan SUPABASE_URL dan SUPABASE_KEY tersedia."
        )
        st.stop()

    return create_client(url, key)
@st.cache_data(ttl=60)
def load_pasar_supabase():
    sb = get_supabase()

    pasar = (
        sb.table("pasar")
        .select("id,nama_pasar,alamat,kecamatan,latitude,longitude,status")
        .execute()
        .data
        or []
    )

    tahunan = (
        sb.table("pasar_tahunan")
        .select("*")
        .execute()
        .data
        or []
    )

    df_pasar = pd.DataFrame(pasar)
    df_tahunan = pd.DataFrame(tahunan)

    if df_pasar.empty or df_tahunan.empty:
        return pd.DataFrame()

    df = df_tahunan.merge(
        df_pasar,
        left_on="pasar_id",
        right_on="id",
        how="left"
    )

    df["lat"] = pd.to_numeric(
        df["latitude"],
        errors="coerce"
    )

    df["lon"] = pd.to_numeric(
        df["longitude"],
        errors="coerce"
    )

    df["tera_ulang_tahun"] = pd.to_numeric(
        df["tahun"],
        errors="coerce"
    )

    df["jumlah_timbangan_tera_ulang"] = pd.to_numeric(
        df["total_uttp"],
        errors="coerce"
    ).fillna(0)
    # =====================================================
    # MAPPING KOLOM SUPABASE -> FORMAT DASHBOARD
    # =====================================================
    
    mapping_timbangan = {
        "timb_pegas": "Timb. Pegas",
        "timb_meja": "Timb. Meja",
        "timb_elektronik": "Timb. Elektronik",
        "timb_sentisimal": "Timb. Sentisimal",
        "timb_bobot_ingsut": "Timb. Bobot Ingsut",
        "neraca": "Neraca",
        "dacin": "Dacin",
    }
    
    for kolom_asal, kolom_dashboard in mapping_timbangan.items():
        if kolom_asal in df.columns:
            df[kolom_dashboard] = pd.to_numeric(
                df[kolom_asal],
                errors="coerce"
            ).fillna(0)
        else:
            df[kolom_dashboard] = 0
    if "total_pedagang" not in df.columns:
        df["total_pedagang"] = 0
    else:
        df["total_pedagang"] = pd.to_numeric(
            df["total_pedagang"],
            errors="coerce"
        ).fillna(0)
    df["nama_pasar"] = df["nama_pasar"].fillna("").astype(str).str.strip()
    df["kecamatan"] = df["kecamatan"].fillna("").astype(str).str.strip()
    df["alamat"] = df["alamat"].fillna("").astype(str).str.strip()
    return df
def build_export_pasar(
    df,
    tahun,
):

    if df.empty:
        return pd.DataFrame()

    tahun = int(tahun)
    tahun_sebelumnya = tahun - 1

    data_tahun = df[
        pd.to_numeric(
            df["tera_ulang_tahun"],
            errors="coerce",
        ) == tahun
    ].copy()

    if data_tahun.empty:
        return pd.DataFrame()

    data_sebelumnya = df[
        pd.to_numeric(
            df["tera_ulang_tahun"],
            errors="coerce",
        ) == tahun_sebelumnya
    ].copy()

    # =====================================================
    # LOOKUP TANGGAL TAHUN SEBELUMNYA
    # =====================================================
    lookup_tanggal_sebelumnya = {}

    if not data_sebelumnya.empty:

        for _, row in data_sebelumnya.iterrows():

            pasar_id = row.get(
                "pasar_id"
            )

            lookup_tanggal_sebelumnya[
                pasar_id
            ] = row.get(
                "tanggal_pelaksanaan"
            )

    hasil = []

    for _, row in data_tahun.iterrows():

        pasar_id = row.get(
            "pasar_id"
        )

        hasil.append(
            {
                "Nama Pasar": row.get(
                    "nama_pasar",
                    "",
                ),

                f"Pelaksanaan {tahun_sebelumnya}": (
                    lookup_tanggal_sebelumnya.get(
                        pasar_id,
                        "",
                    )
                ),

                f"Pelaksanaan {tahun}": row.get(
                    "tanggal_pelaksanaan",
                    "",
                ),

                "TP": int(
                    pd.to_numeric(
                        row.get(
                            "timb_pegas",
                            0,
                        ),
                        errors="coerce",
                    )
                    or 0
                ),

                "TM": int(
                    pd.to_numeric(
                        row.get(
                            "timb_meja",
                            0,
                        ),
                        errors="coerce",
                    )
                    or 0
                ),

                "TE": int(
                    pd.to_numeric(
                        row.get(
                            "timb_elektronik",
                            0,
                        ),
                        errors="coerce",
                    )
                    or 0
                ),

                "Sentisimal": int(
                    pd.to_numeric(
                        row.get(
                            "timb_sentisimal",
                            0,
                        ),
                        errors="coerce",
                    )
                    or 0
                ),

                "Bobot Ingsut": int(
                    pd.to_numeric(
                        row.get(
                            "timb_bobot_ingsut",
                            0,
                        ),
                        errors="coerce",
                    )
                    or 0
                ),

                "Neraca": int(
                    pd.to_numeric(
                        row.get(
                            "neraca",
                            0,
                        ),
                        errors="coerce",
                    )
                    or 0
                ),

                "Total UTTP": int(
                    pd.to_numeric(
                        row.get(
                            "total_uttp",
                            0,
                        ),
                        errors="coerce",
                    )
                    or 0
                ),

                "Total Pedagang": int(
                    pd.to_numeric(
                        row.get(
                            "total_pedagang",
                            0,
                        ),
                        errors="coerce",
                    )
                    or 0
                ),

                "Jumlah Hari": int(
                    pd.to_numeric(
                        row.get(
                            "jumlah_hari",
                            0,
                        ),
                        errors="coerce",
                    )
                    or 0
                ),
            }
        )

    laporan = pd.DataFrame(
        hasil
    )

    if laporan.empty:
        return laporan

    laporan = laporan.sort_values(
        "Nama Pasar"
    ).reset_index(
        drop=True
    )

    laporan.insert(
        0,
        "No",
        range(
            1,
            len(laporan) + 1,
        ),
    )

    return laporan

def generate_excel_rekap_pasar(
    laporan,
    tahun,
):

    output = BytesIO()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Pasar {tahun}"

    tahun_sebelumnya = int(tahun) - 1

    # =====================================================
    # STYLE
    # =====================================================
    thin = Side(
        style="thin",
        color="000000",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    align_center = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    align_left = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )

    font_header = Font(
        name="Arial",
        size=9,
        bold=True,
    )

    font_data = Font(
        name="Arial",
        size=9,
    )

    fill_header = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    fill_total = PatternFill(
        "solid",
        fgColor="E2F0D9",
    )

    # =====================================================
    # JUDUL
    # =====================================================
    ws.merge_cells(
        "A1:M1"
    )

    ws["A1"] = (
        "REKAP DATA UTTP SIDANG TERA ULANG PASAR "
        f"TAHUN {tahun}"
    )

    ws["A1"].font = Font(
        name="Arial",
        size=12,
        bold=True,
    )

    ws["A1"].alignment = align_center

    ws.row_dimensions[1].height = 24

    # =====================================================
    # HEADER
    # =====================================================
    headers = [
        "No",
        "Nama Pasar",
        f"Pelaksanaan {tahun_sebelumnya}",
        f"Pelaksanaan {tahun}",
        "TP",
        "TM",
        "TE",
        "Sentisimal",
        "Bobot Ingsut",
        "Neraca",
        "Total UTTP",
        "Total Pedagang",
        "Jumlah Hari",
    ]

    for col_num, header in enumerate(
        headers,
        start=1,
    ):

        cell = ws.cell(
            row=3,
            column=col_num,
        )

        cell.value = header
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border

    ws.row_dimensions[3].height = 38

    # =====================================================
    # DATA
    # =====================================================
    data_start = 4

    for index, row in laporan.iterrows():

        excel_row = (
            data_start
            + index
        )

        values = [
            row.get("No", ""),
            row.get("Nama Pasar", ""),
            row.get(
                f"Pelaksanaan {tahun_sebelumnya}",
                "",
            ),
            row.get(
                f"Pelaksanaan {tahun}",
                "",
            ),
            row.get("TP", 0),
            row.get("TM", 0),
            row.get("TE", 0),
            row.get("Sentisimal", 0),
            row.get("Bobot Ingsut", 0),
            row.get("Neraca", 0),
            row.get("Total UTTP", 0),
            row.get("Total Pedagang", 0),
            row.get("Jumlah Hari", 0),
        ]

        for col_num, value in enumerate(
            values,
            start=1,
        ):

            # Jangan tampilkan None sebagai tulisan "None"
            if value is None:
                value = ""

            cell = ws.cell(
                row=excel_row,
                column=col_num,
            )

            cell.value = value
            cell.font = font_data
            cell.border = border
            cell.alignment = align_center

        # Nama pasar rata kiri
        ws.cell(
            row=excel_row,
            column=2,
        ).alignment = align_left

        ws.row_dimensions[
            excel_row
        ].height = 25

    # =====================================================
    # TOTAL
    # =====================================================
    total_row = (
        data_start
        + len(laporan)
    )

    ws.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=4,
    )

    total_label = ws.cell(
        row=total_row,
        column=1,
    )

    total_label.value = "JUMLAH"
    total_label.font = font_header
    total_label.alignment = align_center

    # Kolom angka E:M
    kolom_total = [
        "TP",
        "TM",
        "TE",
        "Sentisimal",
        "Bobot Ingsut",
        "Neraca",
        "Total UTTP",
        "Total Pedagang",
        "Jumlah Hari",
    ]

    for offset, nama_col in enumerate(
        kolom_total,
        start=5,
    ):

        nilai = pd.to_numeric(
            laporan[nama_col],
            errors="coerce",
        ).fillna(0).sum()

        ws.cell(
            row=total_row,
            column=offset,
        ).value = int(nilai)

    for col_num in range(
        1,
        14,
    ):

        cell = ws.cell(
            row=total_row,
            column=col_num,
        )

        cell.border = border
        cell.fill = fill_total
        cell.font = font_header
        cell.alignment = align_center

    ws.row_dimensions[
        total_row
    ].height = 24

    # =====================================================
    # LEBAR KOLOM
    # =====================================================
    widths = {
        "A": 6,
        "B": 24,
        "C": 23,
        "D": 23,
        "E": 9,
        "F": 9,
        "G": 9,
        "H": 13,
        "I": 15,
        "J": 10,
        "K": 13,
        "L": 16,
        "M": 12,
    }

    for col, width in widths.items():
        ws.column_dimensions[
            col
        ].width = width

    # =====================================================
    # FREEZE & PRINT
    # =====================================================
    ws.freeze_panes = "E4"

    ws.sheet_view.showGridLines = False

    ws.page_setup.orientation = (
        "landscape"
    )

    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    ws.print_title_rows = "1:3"

    # =====================================================
    # SIMPAN
    # =====================================================
    wb.save(
        output
    )

    output.seek(0)

    return output.getvalue()
@st.cache_data
def load_excel(path_like):
    try:
        if "." not in path_like:
            matches = glob.glob(path_like + ".*")
            path = matches[0] if matches else path_like
        else:
            path = path_like
        df = pd.read_excel(path, engine="openpyxl")
    except:
        st.warning("File Excel tidak ditemukan, menggunakan data sampel.")
        return pd.DataFrame({
            'nama_pasar': ['Cisoka','Curug','Mauk','Cikupa','Pasar Kemis'],
            'kecamatan': ['Cisoka','Curug','Mauk','Cikupa','Pasar Kemis'],
            'alamat': ['Jl. Ps. Cisoka','Jl. Raya Curug','East Mauk','Jl. Raya Serang','RGPJ+FJX'],
            'lat': [-6.26435,-6.26100,-6.06044,-6.22907,-6.16365],
            'lon': [106.42592,106.55858,106.51129,106.51981,106.53155],
            'tera_ulang_tahun': [2025,2025,2025,2025,2025],
            'jumlah_timbangan_tera_ulang': [195,251,161,257,174],
            'jenis_timbangan': ['Pegas:77;Meja:30;Elektronik:87']*5
        }),

    df.columns = [c.strip() for c in df.columns]
    rename = {'Nama Pasar':'nama_pasar','Alamat':'alamat','Kecamatan':'kecamatan',
              'Koordinat':'koordinat','Tahun Tera Ulang':'tera_ulang_tahun',
              'Total UTTP':'jumlah_timbangan_tera_ulang','Total Pedagang':'total_pedagang'}
    df.rename(columns={k:v for k,v in rename.items() if k in df.columns}, inplace=True)

    if 'koordinat' in df.columns:
        coords = df['koordinat'].apply(parse_coord)
        df['lat'] = pd.to_numeric(coords.apply(lambda x: x[0]), errors='coerce')
        df['lon'] = pd.to_numeric(coords.apply(lambda x: x[1]), errors='coerce')

    for col in ['nama_pasar','alamat','kecamatan']:
        if col in df.columns: df[col] = df[col].fillna('').astype(str).str.strip()
    if 'kecamatan' in df.columns: df['kec_norm'] = df['kecamatan'].apply(_norm)
    if 'nama_pasar' in df.columns: df['pasar_norm'] = df['nama_pasar'].apply(_norm)

    return df

@st.cache_data
def load_geojson(path):
    with open(path, 'r', encoding='utf-8') as f:
        gj = json.load(f)
    for ft in gj['features']:
        props = ft['properties']
        wadmkc = props.get('wadmkc','')
        props['kec_norm'] = _norm(wadmkc)
        props['kec_label'] = wadmkc
    return gj


# =========================
# KLIK MARKER -> PENDING
# =========================
def pick_from_click(map_state, df_context, name_col, kec_col, state_prefix):
    if not map_state: return False
    clicked = map_state.get('last_object_clicked')
    if not clicked: return False
    latc, lonc = clicked.get('lat'), clicked.get('lng')
    if None in (latc, lonc): return False
    if not {'lat','lon',name_col,kec_col}.issubset(df_context.columns): return False

    tmp = df_context[['lat','lon',name_col,kec_col]].dropna().copy()
    if tmp.empty: return False
    dist = ((tmp['lat'] - latc)**2 + (tmp['lon'] - lonc)**2).idxmin()
    st.session_state[f"{state_prefix}_pending_pick"] = {
        'name': str(df_context.loc[dist, name_col]),
        'kec': str(df_context.loc[dist, kec_col])
    }
    return True


# =========================
# DASHBOARD PASAR
# =========================
def render_dashboard_pasar():
    df = load_pasar_supabase()
    df_master_pasar = load_master_pasar_supabase()
    geo = load_geojson(FILE_GEOJSON) if os.path.exists(FILE_GEOJSON) else None
    # =====================================================
    # NAVIGASI
    # =====================================================
    col_back, col_home, col_space = st.columns([1.4, 1.4, 5])

    with col_back:
        if st.button(
            "← Dashboard Tera Ulang",
            use_container_width=True,
            key="btn_pasar_kembali_dashboard"
        ):
            st.session_state.halaman_dashboard = "home_dashboard"
            st.rerun()

    with col_home:
        if st.button(
            "🏠 Home SMART METRO",
            use_container_width=True,
            key="btn_pasar_kembali_home"
        ):
            st.session_state.halaman = "home"
            st.session_state.halaman_dashboard = "home_dashboard"
            st.rerun()
    render_main_header(
        "Dashboard Pasar Kabupaten Tangerang",
        "Monitoring tera ulang pasar · Bidang Kemetrologian"
    )
    
    # --- pending click ---
    pending = st.session_state.pop("pasar_pending_pick", None)

    if pending:
        st.session_state['pasar_kec_filter'] = pending['kec']
        st.session_state['pasar_name_filter'] = pending['name']
        st.rerun()

    # --- sidebar filter ---
    st.sidebar.markdown("---"); st.sidebar.subheader("Filter Pasar")
    years = sorted(pd.to_numeric(df['tera_ulang_tahun'], errors='coerce').dropna().astype(int).unique())
    year_pick = st.sidebar.selectbox("Tahun Tera Ulang", years[::-1], key='pasar_year_pick')
    with st.expander(
        "📥 Export Rekap Pasar",
        expanded=False,
    ):

        laporan_pasar = (
            build_export_pasar(
                df=df,
                tahun=year_pick,
            )
        )

        if laporan_pasar.empty:

            st.info(
                f"Belum ada data sidang tera ulang "
                f"pasar tahun {year_pick}."
            )

        else:

            st.success(
                f"Ditemukan "
                f"{len(laporan_pasar)} pasar "
                f"pada tahun {year_pick}."
            )

            st.dataframe(
                laporan_pasar,
                use_container_width=True,
                hide_index=True,
            )
            # =================================================
            # GENERATE FILE EXCEL
            # =================================================
            excel_pasar = generate_excel_rekap_pasar(
                laporan=laporan_pasar,
                tahun=year_pick,
            )
    
            nama_file_pasar = (
                f"Rekap_Data_UTTP_Pasar_"
                f"{year_pick}.xlsx"
            )
    
            st.download_button(
                label="⬇️ Download Excel",
                data=excel_pasar,
                file_name=nama_file_pasar,
                mime=(
                    "application/"
                    "vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                use_container_width=True,
                key="download_rekap_pasar",
            )
    status_pick = st.sidebar.selectbox(
        "Status",
        ["(Semua)", "Sudah Tera", "Belum Tera"],
        key="pasar_status_filter"
    )

    df_year = df[df['tera_ulang_tahun'] == year_pick].copy()
    all_kec = uniq(df_year['kecamatan'], clean=True) if not df_year.empty else []
    all_pasar = uniq(df_year['nama_pasar'], clean=False) if not df_year.empty else []

    kec_pick = st.sidebar.selectbox("Kecamatan", ['(Semua)'] + all_kec,
                                    key='pasar_kec_filter')
    if kec_pick == '(Semua)':
        pasar_ops = ['(Semua)'] + all_pasar
    else:
        pasar_ops = ['(Semua)'] + uniq(df_year[df_year['kecamatan']==kec_pick]['nama_pasar'], clean=False)
    nama_pick = st.sidebar.selectbox("Nama Pasar", pasar_ops, key='pasar_name_filter')

    # simpan state
    st.session_state['pasar_kec_sel'] = kec_pick
    st.session_state['pasar_name_sel'] = nama_pick

    # --- filter dataframe ---
    fdf = df_year.copy()
    if kec_pick != '(Semua)': fdf = fdf[fdf['kecamatan'] == kec_pick]
    if nama_pick != '(Semua)': fdf = fdf[fdf['nama_pasar'] == nama_pick]
    # =====================================================
    # DATA KHUSUS PETA
    # =====================================================
    
    peta_df = df_master_pasar.copy()
    
    if not peta_df.empty:
    
        # Data tahun yang sedang dipilih
        data_tahun = df[
            df["tera_ulang_tahun"] == year_pick
        ].copy()
    
        kolom_status = data_tahun[
            [
                "pasar_id",
                "jumlah_timbangan_tera_ulang"
            ]
        ].copy()
    
        kolom_status["sudah_tera"] = True
    
        peta_df = peta_df.merge(
            kolom_status,
            left_on="id",
            right_on="pasar_id",
            how="left"
        )
    
        peta_df["sudah_tera"] = (
            peta_df["sudah_tera"]
            .fillna(False)
            .astype(bool)
        )
    
        peta_df["jumlah_timbangan_tera_ulang"] = (
            pd.to_numeric(
                peta_df["jumlah_timbangan_tera_ulang"],
                errors="coerce"
            )
            .fillna(0)
        )
    # --- informasi pasar jika spesifik ---
    if nama_pick != '(Semua)' and not fdf.empty:
        r = fdf.iloc[0]
        st.markdown("---")
        st.markdown(f"""
        <div style="background:#f3e8ff; padding:14px 16px; border-radius:12px; border-left:5px solid #8000FF;">
            <h4 style="color:#4B0082;">🏪 {r['nama_pasar']}</h4>
            <p style="font-size:13px;"><b>Kecamatan:</b> {r['kecamatan']}<br>
            <b>Alamat:</b> {r['alamat']}<br><b>Tahun:</b> {year_pick}</p>
        </div>
        """, unsafe_allow_html=True)

    # --- KPI ---
    if nama_pick != '(Semua)':
        cols = st.columns(4)
        cols[0].metric("Nama Pasar", nama_pick)
        cols[1].metric("Kecamatan", fdf['kecamatan'].iloc[0] if not fdf.empty else '-')
        cols[2].metric("Tahun", year_pick)
        cols[3].metric("Total Timbangan", int(fdf['jumlah_timbangan_tera_ulang'].sum()) if not fdf.empty else 0)
    elif kec_pick != '(Semua)':
        cols = st.columns(4)
        cols[0].metric("Kecamatan", kec_pick)
        cols[1].metric("Total Pasar", fdf['nama_pasar'].nunique())
        cols[2].metric("Tahun", year_pick)
        cols[3].metric("Total Timbangan", int(fdf['jumlah_timbangan_tera_ulang'].sum()))
    else:
        total_kecamatan = (
            peta_df["kecamatan"].nunique()
            if not peta_df.empty
            else 0
        )
    
        total_pasar = (
            peta_df["nama_pasar"].nunique()
            if not peta_df.empty
            else 0
        )
    
        sudah_tera = (
            int(peta_df["sudah_tera"].sum())
            if not peta_df.empty and "sudah_tera" in peta_df.columns
            else 0
        )
    
        belum_tera = total_pasar - sudah_tera
    
        total_timbangan = (
            int(fdf["jumlah_timbangan_tera_ulang"].sum())
            if not fdf.empty
            else 0
        )
    
        cols = st.columns(5)
    
        cols[0].metric(
            "Total Kecamatan",
            total_kecamatan
        )
    
        cols[1].metric(
            "Total Pasar",
            total_pasar
        )
    
        cols[2].metric(
            "Sudah Tera",
            sudah_tera
        )
    
        cols[3].metric(
            "Belum Tera",
            belum_tera
        )
    
        cols[4].metric(
            "Total Timbangan",
            total_timbangan
        )

    # =====================================================
    # DAFTAR PASAR BELUM TERA
    # =====================================================
    
    if not peta_df.empty and "sudah_tera" in peta_df.columns:
    
        pasar_belum_tera = peta_df[
            peta_df["sudah_tera"] == False
        ].copy()
    
        if kec_pick != "(Semua)":
            pasar_belum_tera = pasar_belum_tera[
                pasar_belum_tera["kecamatan"] == kec_pick
            ]
    
        if not pasar_belum_tera.empty:
    
            tabel_belum = pasar_belum_tera[
                [
                    "nama_pasar",
                    "kecamatan",
                    "alamat"
                ]
            ].copy()
    
            tabel_belum.rename(
                columns={
                    "nama_pasar": "Nama Pasar",
                    "kecamatan": "Kecamatan",
                    "alamat": "Alamat"
                },
                inplace=True
            )
    
            tabel_belum = tabel_belum.sort_values(
                ["Kecamatan", "Nama Pasar"]
            )
    
            with st.expander(
                f"📋 Pasar Belum Tera ({len(pasar_belum_tera)})",
                expanded=False
            ):
                st.dataframe(
                    tabel_belum,
                    use_container_width=True,
                    hide_index=True
                )
    # --- PETA ---
    st.subheader("🗺️ Peta Lokasi Pasar")
    center, zoom = [-6.2, 106.55], 10
    peta_filter = peta_df.copy()
    if status_pick == "Sudah Tera":
        peta_filter = peta_filter[
            peta_filter["sudah_tera"] == True
        ]
    
    elif status_pick == "Belum Tera":
        peta_filter = peta_filter[
            peta_filter["sudah_tera"] == False
        ]

    if kec_pick != "(Semua)":
        peta_filter = peta_filter[
            peta_filter["kecamatan"] == kec_pick
        ]
    
    if nama_pick != "(Semua)":
        peta_filter = peta_filter[
            peta_filter["nama_pasar"] == nama_pick
        ]
    
    coords = (
        peta_filter[["lat", "lon"]]
        .dropna()
        if {"lat", "lon"}.issubset(peta_filter.columns)
        else pd.DataFrame()
    )

    if not coords.empty:
        if nama_pick != '(Semua)':
            r = fdf[fdf['nama_pasar']==nama_pick].iloc[0]
            center = [float(r['lat']), float(r['lon'])]; zoom = 16
        elif len(coords) == 1:
            center = [coords.iloc[0]['lat'], coords.iloc[0]['lon']]; zoom = 14

    m = folium.Map(location=center, zoom_start=zoom, control_scale=True, tiles=None)
    folium.TileLayer("OpenStreetMap", control=False).add_to(m)
    # =====================================================
    # LEGEND STATUS TERA PASAR
    # =====================================================
    legend_html = f"""
    <div style="
        position: fixed;
        bottom: 35px;
        left: 35px;
        z-index: 9999;
        background-color: white;
        padding: 12px 16px;
        border-radius: 10px;
        border: 1px solid #D1D5DB;
        box-shadow: 0 3px 10px rgba(0,0,0,0.18);
        font-size: 13px;
        min-width: 175px;
    ">

        <div style="
            font-weight:700;
            margin-bottom:8px;
            color:#111827;
        ">
            Status Tera Ulang {year_pick}
        </div>

        <div style="
            display:flex;
            align-items:center;
            gap:8px;
            margin-bottom:6px;
        ">
            <span style="
                width:12px;
                height:12px;
                border-radius:50%;
                background:#16A34A;
                display:inline-block;
            "></span>

            <span>
                Sudah Tera Ulang
            </span>
        </div>

        <div style="
            display:flex;
            align-items:center;
            gap:8px;
        ">
            <span style="
                width:12px;
                height:12px;
                border-radius:50%;
                background:#9CA3AF;
                display:inline-block;
            "></span>

            <span>
                Belum Tera Ulang
            </span>
        </div>

    </div>
    """

    m.get_root().html.add_child(
        folium.Element(
            legend_html
        )
    )
    if geo:
        folium.GeoJson(geo, name="Batas Kecamatan",
                       style_function=lambda x: {"color":"#8000FF","weight":2,"fillOpacity":0},
                       tooltip=folium.GeoJsonTooltip(fields=["kec_label"], aliases=["Kecamatan:"])).add_to(m)

    if not coords.empty:
        cluster = MarkerCluster(name="Pasar").add_to(m)
        for _, r in peta_filter.iterrows():
            if pd.isna(r["lat"]) or pd.isna(r["lon"]):
                continue
        
            if bool(r.get("sudah_tera", False)):
                warna = "#16A34A"
                status_text = "Sudah Tera Ulang"
            else:
                warna = "#9CA3AF"
                status_text = "Belum Tera Ulang"
        
            folium.CircleMarker(
                location=[
                    float(r["lat"]),
                    float(r["lon"])
                ],
                radius=10,
                color=warna,
                fill=True,
                fill_color=warna,
                fill_opacity=0.8,
                weight=2,
                tooltip=r["nama_pasar"],
                popup=folium.Popup(
                    f"""
                    <b>{r['nama_pasar']}</b><br>
                    {r['alamat']}<br>
                    Tahun: {year_pick}<br>
                    Status: {status_text}<br>
                    Total UTTP: {int(r.get('jumlah_timbangan_tera_ulang', 0))}
                    """,
                    max_width=280
                )
            ).add_to(cluster)
        if nama_pick == '(Semua)' and len(coords) > 1:
            m.fit_bounds([[coords['lat'].min(), coords['lon'].min()],
                          [coords['lat'].max(), coords['lon'].max()]], padding=(30,30))

    folium.LayerControl(collapsed=False).add_to(m)
    map_state = st_folium(m, height=500, use_container_width=True, key="pasar_map")

    if pick_from_click(
        map_state,
        peta_filter,
        "nama_pasar",
        "kecamatan",
        "pasar"
    ):
        st.rerun()

    # --- GRAFIK TREN ---
    st.subheader("📈 Grafik (Tahun ke Tahun)")
    gdf = df.copy()
    if nama_pick != '(Semua)': gdf = gdf[gdf['nama_pasar'].str.strip() == nama_pick.strip()]
    elif kec_pick != '(Semua)': gdf = gdf[gdf['kecamatan'] == kec_pick]
    gdf = gdf[pd.to_numeric(gdf['tera_ulang_tahun'], errors='coerce').notna()]
    gdf['tera_ulang_tahun'] = gdf['tera_ulang_tahun'].astype(int)
    # Batasi grafik hanya sampai tahun yang dipilih
    gdf = gdf[
        gdf["tera_ulang_tahun"] <= int(year_pick)
    ].copy()

    agg = gdf.groupby('tera_ulang_tahun').agg(
        jumlah_pasar=('nama_pasar','nunique'),
        total_uttp=('jumlah_timbangan_tera_ulang','sum'),
        total_pedagang=('total_pedagang','sum') if 'total_pedagang' in gdf else ('tera_ulang_tahun','size')
    ).reset_index().sort_values('tera_ulang_tahun')
    agg['Tahun'] = agg['tera_ulang_tahun'].astype(str)

    if agg.empty:
        st.info("Tidak ada data untuk grafik.")
    else:
        import altair as alt
        if (
            kec_pick == "(Semua)"
            and nama_pick == "(Semua)"
        ):

            # =====================================================
            # GRAFIK UMUM KABUPATEN TANGERANG
            # Tampilan disamakan dengan grafik Kecamatan / Pasar
            # =====================================================
            c1, c2 = st.columns(2)

            # =====================================================
            # GRAFIK JUMLAH PEDAGANG
            # =====================================================
            with c1:

                base_pedagang_umum = alt.Chart(
                    agg
                ).encode(
                    x=alt.X(
                        "Tahun:O",
                        title="Tahun",
                        axis=alt.Axis(
                            labelAngle=0
                        )
                    ),
                    y=alt.Y(
                        "total_pedagang:Q",
                        title="Jumlah Pedagang"
                    ),
                    tooltip=[
                        alt.Tooltip(
                            "Tahun:O",
                            title="Tahun"
                        ),
                        alt.Tooltip(
                            "total_pedagang:Q",
                            title="Jumlah Pedagang",
                            format=",.0f"
                        )
                    ]
                )

                area_pedagang_umum = (
                    base_pedagang_umum.mark_area(
                        opacity=0.15
                    )
                )

                line_pedagang_umum = (
                    base_pedagang_umum.mark_line(
                        strokeWidth=3
                    )
                )

                point_pedagang_umum = (
                    base_pedagang_umum.mark_point(
                        filled=True,
                        size=90
                    )
                )

                chart_pedagang_umum = (
                    area_pedagang_umum
                    + line_pedagang_umum
                    + point_pedagang_umum
                ).properties(
                    title=(
                        "Perkembangan Jumlah Pedagang "
                        "- Kabupaten Tangerang"
                    ),
                    height=300
                )

                st.altair_chart(
                    chart_pedagang_umum,
                    use_container_width=True
                )

            # =====================================================
            # GRAFIK JUMLAH TIMBANGAN
            # =====================================================
            with c2:

                base_uttp_umum = alt.Chart(
                    agg
                ).encode(
                    x=alt.X(
                        "Tahun:O",
                        title="Tahun",
                        axis=alt.Axis(
                            labelAngle=0
                        )
                    ),
                    y=alt.Y(
                        "total_uttp:Q",
                        title="Jumlah Timbangan"
                    ),
                    tooltip=[
                        alt.Tooltip(
                            "Tahun:O",
                            title="Tahun"
                        ),
                        alt.Tooltip(
                            "total_uttp:Q",
                            title="Jumlah Timbangan",
                            format=",.0f"
                        )
                    ]
                )

                area_uttp_umum = (
                    base_uttp_umum.mark_area(
                        opacity=0.15
                    )
                )

                line_uttp_umum = (
                    base_uttp_umum.mark_line(
                        strokeWidth=3
                    )
                )

                point_uttp_umum = (
                    base_uttp_umum.mark_point(
                        filled=True,
                        size=90
                    )
                )

                chart_uttp_umum = (
                    area_uttp_umum
                    + line_uttp_umum
                    + point_uttp_umum
                ).properties(
                    title=(
                        "Perkembangan Jumlah Timbangan "
                        "- Kabupaten Tangerang"
                    ),
                    height=300
                )

                st.altair_chart(
                    chart_uttp_umum,
                    use_container_width=True
                )

            # =====================================================
            # GRAFIK JUMLAH PASAR
            # =====================================================
            base_pasar_umum = alt.Chart(
                agg
            ).encode(
                x=alt.X(
                    "Tahun:O",
                    title="Tahun",
                    axis=alt.Axis(
                        labelAngle=0
                    )
                ),
                y=alt.Y(
                    "jumlah_pasar:Q",
                    title="Jumlah Pasar"
                ),
                tooltip=[
                    alt.Tooltip(
                        "Tahun:O",
                        title="Tahun"
                    ),
                    alt.Tooltip(
                        "jumlah_pasar:Q",
                        title="Jumlah Pasar",
                        format=",.0f"
                    )
                ]
            )

            area_pasar_umum = (
                base_pasar_umum.mark_area(
                    opacity=0.15
                )
            )

            line_pasar_umum = (
                base_pasar_umum.mark_line(
                    strokeWidth=3
                )
            )

            point_pasar_umum = (
                base_pasar_umum.mark_point(
                    filled=True,
                    size=90
                )
            )

            chart_pasar_umum = (
                area_pasar_umum
                + line_pasar_umum
                + point_pasar_umum
            ).properties(
                title=(
                    "Perkembangan Jumlah Pasar "
                    "yang Melaksanakan Tera Ulang"
                ),
                height=260
            )

            st.altair_chart(
                chart_pasar_umum,
                use_container_width=True
            )
            # =====================================================
            # PERUBAHAN DIBANDING 1 TAHUN SEBELUMNYA
            # KABUPATEN TANGERANG
            # =====================================================
            tahun_sekarang = int(year_pick)
            tahun_sebelumnya = tahun_sekarang - 1

            data_sekarang = agg[
                agg["tera_ulang_tahun"]
                == tahun_sekarang
            ]

            data_sebelumnya = agg[
                agg["tera_ulang_tahun"]
                == tahun_sebelumnya
            ]

            if not data_sekarang.empty:

                pedagang_sekarang = float(
                    data_sekarang.iloc[0][
                        "total_pedagang"
                    ]
                )

                uttp_sekarang = float(
                    data_sekarang.iloc[0][
                        "total_uttp"
                    ]
                )

                pasar_sekarang = float(
                    data_sekarang.iloc[0][
                        "jumlah_pasar"
                    ]
                )

                if not data_sebelumnya.empty:

                    pedagang_sebelumnya = float(
                        data_sebelumnya.iloc[0][
                            "total_pedagang"
                        ]
                    )

                    uttp_sebelumnya = float(
                        data_sebelumnya.iloc[0][
                            "total_uttp"
                        ]
                    )

                    pasar_sebelumnya = float(
                        data_sebelumnya.iloc[0][
                            "jumlah_pasar"
                        ]
                    )

                    persen_pedagang = (
                        hitung_perubahan_persen(
                            pedagang_sekarang,
                            pedagang_sebelumnya
                        )
                    )

                    persen_uttp = (
                        hitung_perubahan_persen(
                            uttp_sekarang,
                            uttp_sebelumnya
                        )
                    )

                    persen_pasar = (
                        hitung_perubahan_persen(
                            pasar_sekarang,
                            pasar_sebelumnya
                        )
                    )

                else:
                    persen_pedagang = None
                    persen_uttp = None
                    persen_pasar = None

                st.markdown(
                    "#### 📊 Perubahan Kondisi Pasar Kabupaten Tangerang "
                    "dari Tahun Sebelumnya"
                )

                m1, m2, m3 = st.columns(3)

                # =================================================
                # PEDAGANG
                # =================================================
                with m1:

                    if persen_pedagang is not None:

                        arah = (
                            "Naik"
                            if persen_pedagang > 0
                            else "Turun"
                            if persen_pedagang < 0
                            else "Tetap"
                        )

                        tanda = (
                            "+"
                            if persen_pedagang > 0
                            else ""
                        )

                        st.metric(
                            "Jumlah Pedagang",
                            f"{int(pedagang_sekarang):,}"
                            .replace(",", "."),
                            delta=(
                                f"{tanda}"
                                f"{persen_pedagang:.1f}%"
                            )
                        )

                        st.caption(
                            f"{arah} dibanding tahun "
                            f"{tahun_sebelumnya}"
                        )

                    else:
                        st.metric(
                            "Jumlah Pedagang",
                            f"{int(pedagang_sekarang):,}"
                            .replace(",", ".")
                        )

                        st.caption(
                            f"Data tahun "
                            f"{tahun_sebelumnya} "
                            f"belum tersedia."
                        )

                # =================================================
                # TIMBANGAN
                # =================================================
                with m2:

                    if persen_uttp is not None:

                        arah = (
                            "Naik"
                            if persen_uttp > 0
                            else "Turun"
                            if persen_uttp < 0
                            else "Tetap"
                        )

                        tanda = (
                            "+"
                            if persen_uttp > 0
                            else ""
                        )

                        st.metric(
                            "Jumlah Timbangan",
                            f"{int(uttp_sekarang):,}"
                            .replace(",", "."),
                            delta=(
                                f"{tanda}"
                                f"{persen_uttp:.1f}%"
                            )
                        )

                        st.caption(
                            f"{arah} dibanding tahun "
                            f"{tahun_sebelumnya}"
                        )

                    else:
                        st.metric(
                            "Jumlah Timbangan",
                            f"{int(uttp_sekarang):,}"
                            .replace(",", ".")
                        )

                        st.caption(
                            f"Data tahun "
                            f"{tahun_sebelumnya} "
                            f"belum tersedia."
                        )

                # =================================================
                # PASAR
                # =================================================
                with m3:

                    if persen_pasar is not None:

                        arah = (
                            "Naik"
                            if persen_pasar > 0
                            else "Turun"
                            if persen_pasar < 0
                            else "Tetap"
                        )

                        tanda = (
                            "+"
                            if persen_pasar > 0
                            else ""
                        )

                        st.metric(
                            "Jumlah Pasar",
                            f"{int(pasar_sekarang):,}"
                            .replace(",", "."),
                            delta=(
                                f"{tanda}"
                                f"{persen_pasar:.1f}%"
                            )
                        )

                        st.caption(
                            f"{arah} dibanding tahun "
                            f"{tahun_sebelumnya}"
                        )

                    else:
                        st.metric(
                            "Jumlah Pasar",
                            f"{int(pasar_sekarang):,}"
                            .replace(",", ".")
                        )

                        st.caption(
                            f"Data tahun "
                            f"{tahun_sebelumnya} "
                            f"belum tersedia."
                        )
        else:
            c1, c2 = st.columns(2)
        
            # =====================================================
            # JIKA SATU PASAR DIPILIH
            # =====================================================
            if nama_pick != "(Semua)":
        
                with c1:
                    base_pedagang = alt.Chart(agg).encode(
                        x=alt.X(
                            "Tahun:O",
                            title="Tahun",
                            axis=alt.Axis(labelAngle=0)
                        ),
                        y=alt.Y(
                            "total_pedagang:Q",
                            title="Jumlah Pedagang"
                        ),
                        tooltip=[
                            alt.Tooltip("Tahun:O", title="Tahun"),
                            alt.Tooltip(
                                "total_pedagang:Q",
                                title="Jumlah Pedagang",
                                format=",.0f"
                            )
                        ]
                    )
                    
                    area_pedagang = base_pedagang.mark_area(
                        opacity=0.15
                    )
                    
                    line_pedagang = base_pedagang.mark_line(
                        strokeWidth=3
                    )
                    
                    point_pedagang = base_pedagang.mark_point(
                        filled=True,
                        size=90
                    )
                    
                    chart_pedagang = (
                        area_pedagang
                        + line_pedagang
                        + point_pedagang
                    ).properties(
                        title=f"Perkembangan Jumlah Pedagang - {nama_pick}",
                        height=300
                    )
                    st.altair_chart(
                        chart_pedagang,
                        use_container_width=True
                    )
        
                with c2:
                    base_uttp = alt.Chart(agg).encode(
                        x=alt.X(
                            "Tahun:O",
                            title="Tahun",
                            axis=alt.Axis(labelAngle=0)
                        ),
                        y=alt.Y(
                            "total_uttp:Q",
                            title="Jumlah Timbangan"
                        ),
                        tooltip=[
                            alt.Tooltip("Tahun:O", title="Tahun"),
                            alt.Tooltip(
                                "total_uttp:Q",
                                title="Jumlah Timbangan",
                                format=",.0f"
                            )
                        ]
                    )
                    
                    area_uttp = base_uttp.mark_area(
                        opacity=0.15
                    )
                    
                    line_uttp = base_uttp.mark_line(
                        strokeWidth=3
                    )
                    
                    point_uttp = base_uttp.mark_point(
                        filled=True,
                        size=90
                    )
                    
                    chart_uttp = (
                        area_uttp
                        + line_uttp
                        + point_uttp
                    ).properties(
                        title=f"Perkembangan Jumlah Timbangan - {nama_pick}",
                        height=300
                    )
                    st.altair_chart(
                        chart_uttp,
                        use_container_width=True
                    )
                # =====================================================
                # PERUBAHAN DIBANDING 1 TAHUN SEBELUMNYA
                # =====================================================
                
                tahun_sekarang = int(year_pick)
                tahun_sebelumnya = tahun_sekarang - 1
                
                data_sekarang = agg[
                    agg["tera_ulang_tahun"] == tahun_sekarang
                ]
                
                data_sebelumnya = agg[
                    agg["tera_ulang_tahun"] == tahun_sebelumnya
                ]
                
                if not data_sekarang.empty:
                
                    pedagang_sekarang = float(
                        data_sekarang.iloc[0]["total_pedagang"]
                    )
                
                    uttp_sekarang = float(
                        data_sekarang.iloc[0]["total_uttp"]
                    )
                
                    if not data_sebelumnya.empty:
                
                        pedagang_sebelumnya = float(
                            data_sebelumnya.iloc[0]["total_pedagang"]
                        )
                
                        uttp_sebelumnya = float(
                            data_sebelumnya.iloc[0]["total_uttp"]
                        )
                
                        persen_pedagang = hitung_perubahan_persen(
                            pedagang_sekarang,
                            pedagang_sebelumnya
                        )
                
                        persen_uttp = hitung_perubahan_persen(
                            uttp_sekarang,
                            uttp_sebelumnya
                        )
                
                    else:
                        pedagang_sebelumnya = None
                        uttp_sebelumnya = None
                        persen_pedagang = None
                        persen_uttp = None
                st.markdown("#### 📊 Perubahan dari Tahun Sebelumnya")

                c1, c2 = st.columns(2)
            
                with c1:
                    if persen_pedagang is not None:
            
                        arah = "Naik" if persen_pedagang > 0 else (
                            "Turun" if persen_pedagang < 0 else "Tetap"
                        )
            
                        tanda = "+" if persen_pedagang > 0 else ""
            
                        st.metric(
                            "Jumlah Pedagang",
                            f"{int(pedagang_sekarang):,}".replace(",", "."),
                            delta=f"{tanda}{persen_pedagang:.1f}%"
                        )
            
                        st.caption(
                            f"{arah} dibanding tahun {tahun_sebelumnya}"
                        )
            
                    else:
                        st.metric(
                            "Jumlah Pedagang",
                            f"{int(pedagang_sekarang):,}".replace(",", ".")
                        )
            
                        st.caption(
                            f"Data tahun {tahun_sebelumnya} belum tersedia."
                        )
            
                with c2:
                    if persen_uttp is not None:
            
                        arah = "Naik" if persen_uttp > 0 else (
                            "Turun" if persen_uttp < 0 else "Tetap"
                        )
            
                        tanda = "+" if persen_uttp > 0 else ""
            
                        st.metric(
                            "Jumlah Timbangan",
                            f"{int(uttp_sekarang):,}".replace(",", "."),
                            delta=f"{tanda}{persen_uttp:.1f}%"
                        )
            
                        st.caption(
                            f"{arah} dibanding tahun {tahun_sebelumnya}"
                        )
            
                    else:
                        st.metric(
                            "Jumlah Timbangan",
                            f"{int(uttp_sekarang):,}".replace(",", ".")
                        )
            
                        st.caption(
                            f"Data tahun {tahun_sebelumnya} belum tersedia."
                        )
            # =====================================================
            # JIKA HANYA KECAMATAN DIPILIH
            # =====================================================
            else:
            
                # =================================================
                # GRAFIK JUMLAH PEDAGANG
                # =================================================
                with c1:
                    base_pedagang_kec = alt.Chart(agg).encode(
                        x=alt.X(
                            "Tahun:O",
                            title="Tahun",
                            axis=alt.Axis(labelAngle=0)
                        ),
                        y=alt.Y(
                            "total_pedagang:Q",
                            title="Jumlah Pedagang"
                        ),
                        tooltip=[
                            alt.Tooltip(
                                "Tahun:O",
                                title="Tahun"
                            ),
                            alt.Tooltip(
                                "total_pedagang:Q",
                                title="Jumlah Pedagang",
                                format=",.0f"
                            )
                        ]
                    )
            
                    area_pedagang_kec = base_pedagang_kec.mark_area(
                        opacity=0.15
                    )
            
                    line_pedagang_kec = base_pedagang_kec.mark_line(
                        strokeWidth=3
                    )
            
                    point_pedagang_kec = base_pedagang_kec.mark_point(
                        filled=True,
                        size=90
                    )
            
                    chart_pedagang_kec = (
                        area_pedagang_kec
                        + line_pedagang_kec
                        + point_pedagang_kec
                    ).properties(
                        title=f"Perkembangan Jumlah Pedagang - Kecamatan {kec_pick}",
                        height=300
                    )
            
                    st.altair_chart(
                        chart_pedagang_kec,
                        use_container_width=True
                    )
            
                # =================================================
                # GRAFIK JUMLAH TIMBANGAN
                # =================================================
                with c2:
                    base_uttp_kec = alt.Chart(agg).encode(
                        x=alt.X(
                            "Tahun:O",
                            title="Tahun",
                            axis=alt.Axis(labelAngle=0)
                        ),
                        y=alt.Y(
                            "total_uttp:Q",
                            title="Jumlah Timbangan"
                        ),
                        tooltip=[
                            alt.Tooltip(
                                "Tahun:O",
                                title="Tahun"
                            ),
                            alt.Tooltip(
                                "total_uttp:Q",
                                title="Jumlah Timbangan",
                                format=",.0f"
                            )
                        ]
                    )
            
                    area_uttp_kec = base_uttp_kec.mark_area(
                        opacity=0.15
                    )
            
                    line_uttp_kec = base_uttp_kec.mark_line(
                        strokeWidth=3
                    )
            
                    point_uttp_kec = base_uttp_kec.mark_point(
                        filled=True,
                        size=90
                    )
            
                    chart_uttp_kec = (
                        area_uttp_kec
                        + line_uttp_kec
                        + point_uttp_kec
                    ).properties(
                        title=f"Perkembangan Jumlah Timbangan - Kecamatan {kec_pick}",
                        height=300
                    )
            
                    st.altair_chart(
                        chart_uttp_kec,
                        use_container_width=True
                    )
            
                # =====================================================
                # PERUBAHAN DIBANDING 1 TAHUN SEBELUMNYA
                # =====================================================
            
                tahun_sekarang = int(year_pick)
                tahun_sebelumnya = tahun_sekarang - 1
            
                data_sekarang = agg[
                    agg["tera_ulang_tahun"] == tahun_sekarang
                ]
            
                data_sebelumnya = agg[
                    agg["tera_ulang_tahun"] == tahun_sebelumnya
                ]
            
                if not data_sekarang.empty:
            
                    pedagang_sekarang = float(
                        data_sekarang.iloc[0]["total_pedagang"]
                    )
            
                    uttp_sekarang = float(
                        data_sekarang.iloc[0]["total_uttp"]
                    )
            
                    if not data_sebelumnya.empty:
            
                        pedagang_sebelumnya = float(
                            data_sebelumnya.iloc[0]["total_pedagang"]
                        )
            
                        uttp_sebelumnya = float(
                            data_sebelumnya.iloc[0]["total_uttp"]
                        )
            
                        persen_pedagang = hitung_perubahan_persen(
                            pedagang_sekarang,
                            pedagang_sebelumnya
                        )
            
                        persen_uttp = hitung_perubahan_persen(
                            uttp_sekarang,
                            uttp_sebelumnya
                        )
            
                    else:
                        persen_pedagang = None
                        persen_uttp = None
            
                    st.markdown(
                        f"#### 📊 Perubahan Kecamatan {kec_pick} dari Tahun Sebelumnya"
                    )
            
                    m1, m2 = st.columns(2)
            
                    # =================================================
                    # CARD PEDAGANG
                    # =================================================
                    with m1:
                        if persen_pedagang is not None:
            
                            arah = (
                                "Naik"
                                if persen_pedagang > 0
                                else "Turun"
                                if persen_pedagang < 0
                                else "Tetap"
                            )
            
                            tanda = "+" if persen_pedagang > 0 else ""
            
                            st.metric(
                                "Jumlah Pedagang",
                                f"{int(pedagang_sekarang):,}".replace(",", "."),
                                delta=f"{tanda}{persen_pedagang:.1f}%"
                            )
            
                            st.caption(
                                f"{arah} dibanding tahun {tahun_sebelumnya}"
                            )
            
                        else:
                            st.metric(
                                "Jumlah Pedagang",
                                f"{int(pedagang_sekarang):,}".replace(",", ".")
                            )
            
                            st.caption(
                                f"Data tahun {tahun_sebelumnya} belum tersedia."
                            )
            
                    # =================================================
                    # CARD TIMBANGAN
                    # =================================================
                    with m2:
                        if persen_uttp is not None:
            
                            arah = (
                                "Naik"
                                if persen_uttp > 0
                                else "Turun"
                                if persen_uttp < 0
                                else "Tetap"
                            )
            
                            tanda = "+" if persen_uttp > 0 else ""
            
                            st.metric(
                                "Jumlah Timbangan",
                                f"{int(uttp_sekarang):,}".replace(",", "."),
                                delta=f"{tanda}{persen_uttp:.1f}%"
                            )
            
                            st.caption(
                                f"{arah} dibanding tahun {tahun_sebelumnya}"
                            )
            
                        else:
                            st.metric(
                                "Jumlah Timbangan",
                                f"{int(uttp_sekarang):,}".replace(",", ".")
                            )
            
                            st.caption(
                                f"Data tahun {tahun_sebelumnya} belum tersedia."
                            )
                
    # --- TABEL TIMBANGAN (diambil dari fungsi asli) ---
    # (kode tabel timbangan yang panjang tidak diubah, di sini hanya ringkasan)
    # ... (saya akan sisipkan versi singkat, namun di kode asli Anda ada banyak CSS)
    # Agar tetap fungsional, saya sertakan versi sederhana:
        # --- TOTAL TIMBANGAN TERA ULANG (CARD EYECATCHING) ---
    if not fdf.empty and 'jumlah_timbangan_tera_ulang' in fdf.columns:
        st.markdown("---")
        st.subheader("⚖️ Total Timbangan Tera Ulang")

        total_uttp = int(fdf['jumlah_timbangan_tera_ulang'].sum())
        st.markdown(f"""
        <div style="display:flex; justify-content:center;">
            <div style="background:linear-gradient(135deg,#7c3aed,#4c1d95); color:white; 
                        border-radius:16px; padding:20px 40px; box-shadow:0 6px 12px rgba(0,0,0,0.2); 
                        text-align:center; margin-bottom:20px;">
                <div style="font-size:16px; font-weight:600; opacity:0.9;">Total Timbangan Tera Ulang</div>
                <div style="font-size:42px; font-weight:900;">{total_uttp:,}</div>
            </div>
        </div>
        """, unsafe_allow_html=True)

        # Mini card per jenis timbangan (jika kolom tersedia)
        timb_cols = ['Timb. Pegas', 'Timb. Meja', 'Timb. Elektronik', 
                     'Timb. Sentisimal', 'Timb. Bobot Ingsut', 'Neraca', 'Dacin']
        available = [c for c in timb_cols if c in fdf.columns]
        if available:
            st.markdown("""
            <style>
            .mini-card {
                background:white; border-radius:14px; padding:12px; box-shadow:0 3px 6px rgba(0,0,0,0.12);
                border-left:5px solid #7c3aed; margin-bottom:10px;
            }
            .mini-card-title { font-size:13px; font-weight:600; color:#4c1d95; }
            .mini-card-val { font-size:22px; font-weight:800; color:#111827; }
            </style>
            """, unsafe_allow_html=True)

            cols = st.columns(len(available))
            for i, col in enumerate(available):
                val = int(pd.to_numeric(fdf[col], errors='coerce').fillna(0).sum())
                with cols[i]:
                    st.markdown(f"""
                    <div class="mini-card">
                        <div class="mini-card-title">{col.replace('Timb. ','')}</div>
                        <div class="mini-card-val">{val:,}</div>
                    </div>
                    """, unsafe_allow_html=True)


# =========================================================
# ENTRY POINT DASHBOARD PASAR
# =========================================================
def run():
    render_dashboard_pasar()
